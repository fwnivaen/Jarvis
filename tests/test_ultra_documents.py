import importlib.util
import tempfile
import unittest
import zipfile
from pathlib import Path
import ultra_documents
import base64
from dataclasses import replace
from unittest.mock import patch
import server
import ultra
import build
import note_store
from test_offline import request


def write_pdf(path):
    content = b"BT /F1 12 Tf 50 700 Td (Quarterly budget is 120 pounds.) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>", b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"]
    data, offsets = b"%PDF-1.4\n", [0]
    for n, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += str(n).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    xref = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n" + b"".join(("%010d 00000 n \n" % off).encode() for off in offsets[1:])
    data += b"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n" + str(xref).encode() + b"\n%%EOF\n"
    path.write_bytes(data)


class DocumentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.notes = self.root / "notes"
        self.notes.mkdir()

    def test_filename_and_source_links_are_rejected(self):
        for name in ("../../config.json", "../a.pdf", "/a.pdf", "a.exe", "a\\b.pdf"):
            with self.assertRaises(ultra_documents.DocumentError):
                ultra_documents.filename(name)
        target = self.root / "actual.pdf"
        write_pdf(target)
        link = self.root / "linked.pdf"
        link.symlink_to(target)
        with self.assertRaises(ultra_documents.DocumentError):
            ultra_documents.convert(link)

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "Ultra venv optional dependency")
    def test_pdf_text_becomes_idempotent_markdown_with_provenance(self):
        path = self.root / "budget.pdf"
        write_pdf(path)
        first = ultra_documents.import_document(path, self.notes)
        self.assertIn("120 pounds", first["text"])
        self.assertIn("SHA-256:", first["text"])
        self.assertTrue(first["created"])
        reloaded = build.read_notes(self.notes)[0]
        self.assertEqual(reloaded["label"], "budget")
        self.assertEqual(note_store.read(reloaded, self.notes)["title"], "budget")
        self.assertFalse(ultra_documents.import_document(path, self.notes)["created"])
        (self.notes / first["path"]).write_text("Edited by user")
        with self.assertRaises(ultra_documents.DocumentError):
            ultra_documents.import_document(path, self.notes)

    @unittest.skipUnless(importlib.util.find_spec("pypdf"), "Ultra venv optional dependency")
    def test_http_import_publishes_note_without_an_approval_gate(self):
        path = self.root / "budget.pdf"
        write_pdf(path)
        body = {"filename": path.name, "data": base64.b64encode(path.read_bytes()).decode("ascii")}
        settings = replace(ultra.Settings(), enabled=True)
        with patch.object(server, "ultra_settings", return_value=settings), patch.object(server, "NOTES", []), \
                patch.object(server.build, "NOTES_DIR", self.notes), patch.object(server.build, "write_graph"), \
                patch.object(server, "INBOX", self.root / "inbox"):
            status, _, _ = request("/memory/import", "POST", body, headers={"Origin": "http://127.0.0.1:4700"})
            self.assertEqual(status, 200)
            self.assertEqual(len(server.NOTES), 1)
            self.assertIn("120 pounds", server.NOTES[0]["text"])
            self.assertEqual(request("/memory/import", "POST", body, headers={"Origin": "http://127.0.0.1:4700"})[0], 200)
            self.assertEqual(len(server.NOTES), 1)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "Ultra venv optional dependency")
    def test_docx_preserves_paragraph_and_table_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Before table")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Value"
        table.cell(1, 0).text = "Budget"
        table.cell(1, 1).text = "120"
        document.add_paragraph("After table")
        path = self.root / "budget.docx"
        document.save(path)
        text = ultra_documents.convert(path)
        self.assertLess(text.index("Before table"), text.index("Budget"))
        self.assertLess(text.index("Budget"), text.index("After table"))

    @unittest.skipUnless(importlib.util.find_spec("openpyxl"), "Ultra venv optional dependency")
    def test_spreadsheet_preserves_formulas_without_executing_them(self):
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.active.append(["Budget", "Tax"])
        workbook.active.append([120, "=A2*0.2"])
        path = self.root / "budget.xlsx"
        workbook.save(path)
        text = ultra_documents.convert(path)
        self.assertIn("=A2*0.2", text)
        self.assertIn("120", text)
        self.assertIn("no formulas or macros were executed", text)


if __name__ == "__main__":
    unittest.main()
