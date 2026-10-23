"""Bounded, explicit PDF/DOCX/XLSX imports into the existing Markdown vault."""
import hashlib
import os
import re
import time
import zipfile
from pathlib import Path

MAX_DOCUMENT_BYTES = 16 * 1024 * 1024
MAX_EXPANDED_BYTES = 80 * 1024 * 1024
MAX_ZIP_ENTRIES = 2000
MAX_PAGES = 200
MAX_SHEETS = 40
MAX_ROWS = 10000
MAX_COLUMNS = 100
MAX_TEXT_CHARS = 200000
MAX_FILENAME_CHARS = 160
HASH_CHARS = 16
MAX_SLUG_CHARS = 80
PRIVATE_FILE_MODE = 0o600
SUFFIXES = frozenset({".pdf", ".docx", ".xlsx"})


class DocumentError(ValueError):
    pass


def filename(value):
    if (not isinstance(value, str) or not 0 < len(value) <= MAX_FILENAME_CHARS
            or Path(value).name != value or "/" in value or "\\" in value
            or Path(value).suffix.lower() not in SUFFIXES):
        raise DocumentError("Choose a PDF, DOCX or XLSX file with a simple filename.")
    return value


def _cell(value):
    return str(value if value is not None else "").replace("|", "\\|").replace("\n", "<br>")


def _table(rows):
    rows = list(rows)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    rendered = ["| " + " | ".join(_cell(cell) for cell in list(row) + [""] * (width - len(row))) + " |" for row in rows]
    return "\n".join([rendered[0], "| " + " | ".join(["---"] * width) + " |"] + rendered[1:])


def convert(path):
    path = Path(path)
    filename(path.name)
    if path.is_symlink() or not path.is_file() or not 0 < path.stat().st_size <= MAX_DOCUMENT_BYTES:
        raise DocumentError("The document is missing, linked or exceeds 16 MiB.")
    suffix = path.suffix.lower()
    try:
        if suffix != ".pdf":
            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_ZIP_ENTRIES or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                    raise DocumentError("The expanded document exceeds the import limit.")
        if suffix == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted or len(reader.pages) > MAX_PAGES:
                raise DocumentError("Use an unlocked PDF with at most 200 pages.")
            parts = []
            count, has_text = 0, False
            for number, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                count += len(text)
                has_text |= bool(text.strip())
                if count > MAX_TEXT_CHARS:
                    raise DocumentError("The extracted document exceeds 200,000 characters.")
                parts.append("## Page " + str(number) + "\n\n" + text)
            if not has_text:
                raise DocumentError("This PDF has no extractable text. Run OCR before importing it.")
            text = "\n\n".join(parts)
        elif suffix == ".docx":
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            document = Document(path)
            parts = []
            for block in document.iter_inner_content():
                if isinstance(block, Paragraph):
                    parts.append(block.text)
                elif isinstance(block, Table):
                    parts.append(_table([[cell.text for cell in row.cells] for row in block.rows]))
                if sum(map(len, parts)) > MAX_TEXT_CHARS:
                    raise DocumentError("The extracted document exceeds 200,000 characters.")
            text = "\n\n".join(parts)
        else:
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                if len(workbook.worksheets) > MAX_SHEETS:
                    raise DocumentError("The workbook exceeds 40 sheets.")
                parts, total_rows, total_chars = [], 0, 0
                for sheet in workbook.worksheets:
                    rows = []
                    for row in sheet.iter_rows(values_only=True):
                        total_rows += 1
                        if total_rows > MAX_ROWS or len(row) > MAX_COLUMNS:
                            raise DocumentError("The workbook exceeds 10,000 rows or 100 columns.")
                        if any(cell is not None for cell in row):
                            rows.append(row)
                            total_chars += sum(len(str(cell)) for cell in row if cell is not None)
                        if total_chars > MAX_TEXT_CHARS:
                            raise DocumentError("The extracted document exceeds 200,000 characters.")
                    parts.append("## " + sheet.title + "\n\n" + _table(rows))
                text = "Formulas are preserved as text; no formulas or macros were executed.\n\n" + "\n\n".join(parts)
            finally:
                workbook.close()
        if not text.strip() or len(text) > MAX_TEXT_CHARS:
            raise DocumentError("The extracted document is empty or exceeds 200,000 characters.")
        return text
    except ImportError:
        raise DocumentError("Document import needs the Ultra virtual environment. Run tools/setup_ultra.py.") from None
    except DocumentError:
        raise
    except Exception:
        raise DocumentError("The document could not be read. Check its format and encryption.") from None


def import_document(path, notes_dir, deadline=None):
    path, root = Path(path), Path(notes_dir).resolve()
    text = convert(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    title = Path(path.name).stem
    slug = re.sub(r"[^\w -]", "", title).strip()[:MAX_SLUG_CHARS] or "Document"
    folder = root / "imports"
    if folder.is_symlink() or not root.is_dir():
        raise DocumentError("The notes import folder is unavailable.")
    folder.mkdir(exist_ok=True)
    target = folder / (slug + "-" + digest[:HASH_CHARS] + ".md")
    content = "# " + title + "\n\nSource file: " + path.name + "\nSHA-256: " + digest + "\n\n" + text + "\n"
    if deadline is not None and time.monotonic() >= deadline:
        raise DocumentError("Import time budget exhausted before writing a note.")
    if target.is_symlink():
        raise DocumentError("The import destination is linked.")
    created = False
    try:
        descriptor = os.open(target, os.O_WRONLY | os.O_CREAT | os.O_EXCL, PRIVATE_FILE_MODE)
        with os.fdopen(descriptor, "w", encoding="utf-8") as handle:
            handle.write(content)
            handle.flush()
            os.fsync(handle.fileno())
        created = True
    except FileExistsError:
        if target.read_text(encoding="utf-8") != content:
            raise DocumentError("An existing import was edited; it will not be overwritten.") from None
    return {"path": target.relative_to(root).as_posix(), "title": title, "text": content,
        "created": created, "sha256": digest}
